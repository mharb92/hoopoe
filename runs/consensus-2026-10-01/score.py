"""Score the consensus run (chat 34) per src/scoring.md, including its chat 34 sections.

Inputs (all in this folder): answer-key.json, control-ids.json, results/<runner>/<model>/batch-*.json.
Outputs: scored/. Nothing is written to the database.

    python3 score.py                  controls first; stops with exit 2 if the 70% gate fails
    python3 score.py --gate-decided   score the queue anyway, once the gate decision is recorded in src/scoring.md
"""
import csv, json, re, sys, unicodedata
from collections import Counter
from pathlib import Path

RUN = Path(__file__).resolve().parent
OUT = RUN / 'scored'
SCORED = ['omar/opus-5.5', 'marwan/gpt-6.1-sol', 'omar/gemini-3.8-flash']   # newest per maker
INDEPENDENT = (1, 2)   # SCORED positions not from the maker whose model produced our values (Opus 5.5)
EVIDENCE = ['omar/opus-5', 'omar/opus-4.8', 'marwan/sonnet-5.5', 'omar/gpt-5.6-sol', 'omar/gpt-5.6-terra']
FIELDS = ['romanization', 'vowelled_arabic', 'pronunciation_confidence', 'meaning_ok', 'meaning_fix',
          'dialect', 'palestinian_alternative', 'reasoning']
GATE_AB = 0.70
GATE_C = 0.95
RAISED = ('A', 'A2', 'B')

MARKS = re.compile(r'[\u064B-\u0650\u0652\u0653-\u065F]')   # vowels, tanween, sukun, other marks; not shadda
NOT_ARABIC = re.compile(r'[^\u0621-\u064A\u0651\u064B-\u0652 ]')   # punctuation and anything else
REGIONAL = re.compile(r'galilee|gaza|hebron|khalil|nablus|jenin|jaffa|haifa|village|rural|fellah|bedouin|regional|northern|southern',
                      re.I)
OK_DIALECT = {'palestinian', 'pan_levantine'}
DIGRAPHS = {'sh': '1', 'gh': '2', 'kh': '4', 'th': '5', 'dh': '6', 'TH': '8'}   # one symbol per sound while matching
CONS = r"[b-df-hj-np-tv-zDST'3712456 8]".replace(' ', '')


def norm(s):
    return re.sub(r'\s+', ' ', unicodedata.normalize('NFC', s or '')).strip()


def sound(s):
    """Romanization match key: forgives notation, short e/i and o/u, and helping vowels; keeps every other contrast."""
    s = norm(s)
    for a, b in DIGRAPHS.items():
        s = s.replace(a, b)
    s = re.sub(r"(?<![a-zA-Z'0-9])(?:a|i)?llaa?h?(?![a-zA-Z'0-9])", 'ALLAH', s)            # allah / allaah / alla
    s = re.sub(r"(?<![a-zA-Z'0-9])(?:w|u|wu|wi)(?:-| )", 'w ', s)                          # and: w / u- / wi-
    s = re.sub(r"(?<=[aeiou]) ?i?(l|[b-df-hj-np-tv-zDST'3125]{1,2})-", r'\1-', s)        # the: elided after a vowel
    s = re.sub(r"(?<![a-zA-Z'0-9])i?(l|[b-df-hj-np-tv-zDST'3125]{1,2})-", 'il-', s)      # the: il- and assimilated forms
    s = re.sub(rf'(?<=[aeiou]{CONS})[iueo](?={CONS}(?:$|[ -]))', '', s)                   # helping vowel: CVCiC = CVCC
    s = re.sub(rf"(?<![a-zA-Z'0-9])([ybtn])i(?={CONS}[aeiou])", r'\1', s)                 # yi- bi- ti- ni- = y- b- t- n-
    s = s.replace('-', '').replace(' ', '')
    s = re.sub(r'(?<!e)e(?!e)', 'i', s)                                                     # short e = i
    return re.sub(r'(?<!o)o(?!o)', 'u', s)                                                  # short o = u


def _arabic(s):
    s = norm(s).replace('\u0649\u0670', '\u0649').replace('\u0670', '\u0627')
    return re.sub(r'\s+', ' ', NOT_ARABIC.sub('', s.replace('\u0652', '').replace('\u0640', ''))).strip()


def letters(s):
    """Letters and shadda only: the vowelled-Arabic match key."""
    return MARKS.sub('', _arabic(s))


def _marks_key(s):
    s = _arabic(s)
    s = re.sub('\u064E(?=\u064A(?![\u064B-\u0651]))', '\u0650', s)   # fatha+ya (ee) equals kasra+ya
    s = re.sub('\u064E(?=\u0648(?![\u064B-\u0651]))', '\u064F', s)   # fatha+waw (oo) equals damma+waw
    # one cluster per letter, its marks sorted, so shadda+vowel order does not matter
    return [c[0] + ''.join(sorted(c[1:])) for c in re.findall(r'.[\u064B-\u0651]*', s)]


def marks_differ(a, b):
    return _marks_key(a) != _marks_key(b)


def rom_lvq(a, b):
    """True when a and b differ only in long-vowel quality (ee/ii, oo/uu)."""
    a, b = sound(a), sound(b)
    q = lambda s: s.replace('ee', 'ii').replace('oo', 'uu')
    return a != b and q(a) == q(b)


def msa_labelled(english):
    return bool(re.search(r'\bMSA\b|formal', english, re.I))


def flags(key, runs):
    out = []
    for n, r in enumerate(runs):
        if r['meaning_ok'] != 'yes':
            out.append(f'{n}:meaning {r["meaning_ok"]}')
        d = r['dialect']
        if d == 'other_levantine' or (d == 'msa' and not msa_labelled(key['english'])):
            out.append(f'{n}:dialect {d}')
        if norm(r['palestinian_alternative']):
            out.append(f'{n}:alternative')
    return out


def dialect_runs(fl):
    """How many runs say the word is not Palestinian or offer another word."""
    return len({f[0] for f in fl if 'meaning' not in f})


def _match(r, key):
    return sound(r['romanization']) == sound(key['our_romanization']) and \
        letters(r['vowelled_arabic']) == letters(key['our_vowelled'])


def outcome(key, runs):
    """Outcome for one word from the three scored runs (src/scoring.md)."""
    ours_r, ours_l = sound(key['our_romanization']), letters(key['our_vowelled'])
    rom = [sound(r['romanization']) for r in runs]
    let = [letters(r['vowelled_arabic']) for r in runs]
    exact = sum(x == ours_r for x in rom)
    lvq = sum(rom_lvq(r['romanization'], key['our_romanization']) for r in runs)
    vow_ok = all(x == ours_l for x in let)
    unanimous = len(set(rom)) == 1 and len(set(let)) == 1
    fl = flags(key, runs)
    clean = not fl and all(r['dialect'] in OK_DIALECT or (r['dialect'] == 'msa' and msa_labelled(key['english']))
                           for r in runs) and all(r['pronunciation_confidence'] != 'low' for r in runs)
    if vow_ok and clean and exact == 3:
        o = 'A'
    elif vow_ok and clean and exact == 2 and lvq == 1:
        o = 'B'
    elif clean and all(_match(runs[n], key) for n in INDEPENDENT):
        o = 'A2'
    elif fl:
        o = 'D2' if dialect_runs(fl) >= 2 else 'D'
    elif any(r['pronunciation_confidence'] == 'low' for r in runs):
        o = 'E'
    elif unanimous and (rom[0] != ours_r or let[0] != ours_l):
        o = 'C'
    else:
        o = 'E'
    return {'outcome': o, 'exact': exact, 'lvq': lvq, 'vow_match': sum(x == ours_l for x in let),
            'unanimous': unanimous, 'unanimous_matches_ours': unanimous and rom[0] == ours_r and let[0] == ours_l,
            'short_vowel_diffs': sum(marks_differ(r['vowelled_arabic'], key['our_vowelled']) for r in runs),
            'flags': fl}


def proposal(key, runs):
    """The pre-filled answer for a review row: a form two runs agree on, else ours."""
    ours = (sound(key['our_romanization']), letters(key['our_vowelled']))
    groups = Counter((sound(r['romanization']), letters(r['vowelled_arabic'])) for r in runs)
    best, n = max(((g, c) for g, c in groups.items() if g != ours), key=lambda x: x[1], default=(None, 0))
    if n >= 2:
        members = [r for r in runs if (sound(r['romanization']), letters(r['vowelled_arabic'])) == best]
        vows = [norm(r['vowelled_arabic']) for r in members]
        return {'romanization': norm(members[0]['romanization']), 'vowelled': max(vows, key=vows.count),
                'why': f'{n} of 3 models agree on this form, not ours'}
    m = groups.get(ours, 0)
    why = (f'{m} of 3 models match ours' if m >= 2 else
           'no majority: ours shown, 1 model matches it' if m == 1 else 'no majority: ours shown, no model matches it')
    return {'romanization': key['our_romanization'], 'vowelled': key['our_vowelled'], 'why': why}


def _consonants(s, merge=False):
    """Consonants only; with merge, also q = ', th = t, dh = d, doubled = single."""
    s = sound(s)
    if merge:
        s = s.replace('q', "'").replace('5', 't').replace('6', 'd').replace('8', 'D')
    s = re.sub(r'[aeiou]', '', s)
    return re.sub(r'(.)\1', r'\1', s) if merge else s


ISSUES = ['', 'vowels', "q or ', th or t, or doubling", 'consonants']


def issue(key, runs, res):
    """Plain-words label for what a review row disagrees on."""
    if dialect_runs(res['flags']):
        return 'one model says not Palestinian'
    worst = ''
    for r in runs:
        a, b = key['our_romanization'], r['romanization']
        if sound(a) == sound(b):
            continue
        kind = ('vowels' if _consonants(a) == _consonants(b) else
                ISSUES[2] if _consonants(a, True) == _consonants(b, True) else 'consonants')
        worst = max(worst, kind, key=ISSUES.index)
    if not worst:
        worst = 'Arabic letters only' if res['vow_match'] < 3 else 'a model is unsure'
    return worst + (' + meaning question' if res['flags'] else '')


def control_rates(rows):
    n = len(rows)
    un = [r for r in rows if r['unanimous']]
    return {'n': n, 'ab': sum(r['outcome'] in RAISED for r in rows) / n,
            'unanimous': len(un), 'unanimous_match': (sum(r['unanimous_matches_ours'] for r in un) / len(un)) if un else 0.0}


def load_run(name, ids):
    got = {}
    for f in sorted((RUN / 'results' / name).glob('batch-*.json')):
        for e in json.loads(f.read_text(encoding='utf-8')):
            got[str(e['id'])] = e
    missing = ids - got.keys()
    if missing:
        sys.exit(f'{name}: {len(missing)} ids missing, e.g. {sorted(missing)[:5]}')
    return got


def why(res, names):
    parts = [f'romanization {res["exact"]}/3 match' + (f', {res["lvq"]} ee/ii-oo/uu only' if res['lvq'] else ''),
             f'vowelled letters {res["vow_match"]}/3']
    if res['short_vowel_diffs']:
        parts.append(f'short vowels differ in {res["short_vowel_diffs"]}/3')
    if res['flags']:
        parts.append('flags: ' + '; '.join(f'{names[int(f[0])]} {f[2:]}' for f in res['flags']))
    if res['unanimous'] and not res['unanimous_matches_ours']:
        parts.append('the three agree with each other, not with us')
    return '. '.join(parts)


def write_csv(path, rows, cols):
    with open(path, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction='ignore')
        w.writeheader()
        w.writerows({k: ('yes' if v is True else '' if v is False else v) for k, v in r.items()} for r in rows)


def main():
    key = json.loads((RUN / 'answer-key.json').read_text(encoding='utf-8'))
    controls = {str(i) for i in json.loads((RUN / 'control-ids.json').read_text())}
    ids = set(key)
    runs = {n: load_run(n, ids) for n in SCORED + EVIDENCE}
    short = {n: n.split('/')[1] for n in runs}
    OUT.mkdir(exist_ok=True)

    rows = []
    for i in sorted(ids, key=int):
        k = key[i]
        scored = [runs[n][i] for n in SCORED]
        res = outcome(k, scored)
        all_rom = {sound(runs[n][i]['romanization']) for n in runs} | {sound(k['our_romanization'])}
        regional = any(REGIONAL.search(runs[n][i]['reasoning'] or '') for n in runs) and bool(res['flags'])
        prop = proposal(k, scored)
        row = {'id': int(i), 'control': i in controls, 'level': k['level'], 'outcome': res['outcome'],
               'why': why(res, [short[n] for n in SCORED]), 'disagreement': len(all_rom),
               'possible_regional': 'yes' if regional else '', 'arabic': k['arabic'], 'english': k['english'],
               'pos': k['pos'], 'our_romanization': k['our_romanization'], 'our_vowelled': k['our_vowelled'],
               'confidence_now': k['confidence_now'], 'hold_reasons': k['hold_reasons'],
               'issue': issue(k, scored, res), 'proposed_romanization': prop['romanization'],
               'proposed_vowelled': prop['vowelled'], 'why_proposed': prop['why'],
               'other_word_offered': ' / '.join(dict.fromkeys(norm(r['palestinian_alternative']) for r in scored
                                                              if norm(r['palestinian_alternative']))),
               'dialect_labels': ', '.join(r['dialect'] for r in scored),
               'meaning_fixes': ' / '.join(dict.fromkeys(norm(r['meaning_fix']) for r in scored if norm(r['meaning_fix']))),
               'meaning_votes': sum(r['meaning_ok'] != 'yes' for r in scored), **res}
        for n in SCORED + EVIDENCE:
            tag = short[n] + ('*' if n in SCORED else '')
            for fld in FIELDS:
                row[f'{tag} {fld}'] = runs[n][i].get(fld, '')
        rows.append(row)

    run_cols = [f'{short[n]}{"*" if n in SCORED else ""} {fld}' for n in SCORED + EVIDENCE for fld in FIELDS]
    base = ['id', 'level', 'outcome', 'result', 'why', 'disagreement', 'possible_regional', 'arabic', 'english', 'pos',
            'our_romanization', 'our_vowelled', 'confidence_now', 'hold_reasons']
    ctrl = [r for r in rows if r['control']]
    queue = [r for r in rows if not r['control']]
    rates = control_rates(ctrl)
    count = lambda rs: dict(Counter(r['outcome'] for r in rs).most_common())
    write_csv(OUT / 'controls.csv', ctrl, base + run_cols)

    report = ['# Consensus scoring\n',
              f'Scored runs: {", ".join(short[n] for n in SCORED)}. Evidence only: {", ".join(short[n] for n in EVIDENCE)}. '
              'Romanization matched by sound (src/scoring.md, chat 34 second section).\n',
              f'## Controls ({rates["n"]} trusted words)\n',
              f'Outcomes: {count(ctrl)}. **Raised (A, A2, B): {rates["ab"]:.1%}** (gate {GATE_AB:.0%}). '
              f'Three runs unanimous on {rates["unanimous"]}; of those, {rates["unanimous_match"]:.1%} match ours '
              f'(automatic adoption needs {GATE_C:.0%}).\n']
    unan_wrong = [r for r in ctrl if r['unanimous'] and not r['unanimous_matches_ours']]
    if unan_wrong:
        report.append('Controls the three runs unanimously disagree with (each needs a look):\n')
        report += [f'- {r["id"]} {r["arabic"]} "{r["english"]}": ours `{r["our_romanization"]}` {r["our_vowelled"]}, '
                   f'runs `{r["proposed_romanization"]}` {r["proposed_vowelled"]}' for r in unan_wrong]
        report.append('')
    if rates['ab'] < GATE_AB:
        report.append(f'**Gate failed: {rates["ab"]:.1%} of controls raised, under {GATE_AB:.0%}.**\n')
        report += _diagnostics(ctrl)
        if '--gate-decided' not in sys.argv:
            report.append('The queue is not scored until the thresholds are revisited (src/scoring.md).\n')
            (OUT / 'report.md').write_text('\n'.join(report), encoding='utf-8')
            print('\n'.join(report))
            sys.exit(2)
        report.append('Decision recorded in src/scoring.md; the queue is scored.\n')

    c_auto = rates['unanimous_match'] >= GATE_C
    english_only = lambda r: r['flags'] and not dialect_runs(r['flags']) and r['why_proposed'] in (
        '2 of 3 models match ours', '3 of 3 models match ours')
    for r in queue + ctrl:
        o = r['outcome']
        if r['control'] and o not in ('D', 'D2'):
            r['result'] = 'trusted, unchanged'
        elif o in RAISED:
            r['result'] = 'raise to 3'
        elif o == 'C':
            r['result'] = 'adopt alternative' if c_auto else 'Arabic review'
        elif o == 'D2':
            r['result'] = 'not taught'
        elif english_only(r):
            r['result'] = 'English review'
        else:
            r['result'] = 'Arabic review'
    order = lambda rs: sorted(rs, key=lambda r: (r['level'], -r['disagreement'], r['id']))
    pick = lambda res: order([r for r in queue + ctrl if r['result'] == res])
    raised, adopted = pick('raise to 3'), pick('adopt alternative')
    arabic, english, not_taught = pick('Arabic review'), pick('English review'), pick('not taught')
    for r in adopted:
        r['new_romanization'], r['new_vowelled'] = r['proposed_romanization'], r['proposed_vowelled']

    write_csv(OUT / 'raised.csv', raised, ['id', 'level', 'outcome', 'arabic', 'english', 'our_romanization', 'our_vowelled', 'why'])
    write_csv(OUT / 'adopted.csv', adopted, base + ['new_romanization', 'new_vowelled'] + run_cols[:3 * len(FIELDS)])
    sheets = _sheets(run_cols)
    for name, data in (('arabic', arabic), ('english', english), ('not-taught', not_taught)):
        write_csv(OUT / f'review-{name}.csv', data, sheets[name])
    _workbook(OUT / 'review.xlsx', {'Arabic': (arabic, sheets['arabic']), 'English meaning': (english, sheets['english']),
                                    'Not taught': (not_taught, sheets['not-taught'])})

    review_total = len(arabic) + len(english) + len(not_taught)
    report += [f'## Queue ({len(queue)} words)\n', f'Outcomes: {count(queue)}.\n',
               f'- Raised to 3: {sum(not r["control"] for r in raised)} (A and B: all three agree with ours; '
               f'A2: GPT and Gemini both match ours exactly)',
               f'- Alternatives adopted automatically: {len(adopted)}',
               f'- Review, {review_total} rows in `review.xlsx`, trusted controls included: Arabic {len(arabic)}, '
               f'English meaning {len(english)}, not taught by rule {len(not_taught)} (2-3 models say not Palestinian)\n']
    report += _diagnostics(queue)
    (OUT / 'report.md').write_text('\n'.join(report), encoding='utf-8')
    print('\n'.join(report))


def _sheets(run_cols):
    ev = [c for c in run_cols if c.split(' ', 1)[1] in ('romanization', 'vowelled_arabic', 'pronunciation_confidence',
                                                        'dialect', 'palestinian_alternative', 'meaning_fix', 'reasoning')]
    head = ['id', 'level', 'arabic', 'english', 'our_romanization', 'our_vowelled']
    return {
        'arabic': head + ['issue', 'proposed_romanization', 'proposed_vowelled', 'why_proposed', 'decision', 'notes',
                          'other_word_offered', 'meaning_fixes', 'possible_regional', 'control'] + ev,
        'english': head + ['meaning_votes', 'meaning_fixes', 'decision', 'notes', 'control'] + ev,
        'not-taught': head + ['other_word_offered', 'dialect_labels', 'decision', 'notes', 'possible_regional', 'control'] + ev,
    }


INSTRUCTIONS = [
    'How to review',
    '',
    'Arabic tab: each row shows our romanization and vowelled Arabic, and a proposed answer.',
    '  The proposal is the form two of the three scored models agree on; if none agree, it is ours.',
    '  decision: leave empty to accept the proposal. Write "ours" to keep ours. Otherwise write the right',
    '  romanization (and Arabic, if it differs) in notes and put "other" in decision.',
    '  Write "drop" if the word should not be taught at all.',
    'English meaning tab: the Arabic is agreed; only the English gloss is questioned.',
    '  decision: leave empty to keep our English. Otherwise write the better English in notes and "fix" in decision.',
    'Not taught tab: two or three models say the word is not Palestinian. It will not be taught.',
    '  decision: leave empty to agree. Write "teach" if Palestinians do say it.',
    '',
    'Rows run from the commonest words (level 1) down, most disputed first within a level.',
    'Columns after "control" are every model\'s answer and reasoning; a * marks the three scored models.',
    'control = yes marks a word we already trusted that the models flagged.',
]


def _workbook(path, tabs):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'How to review'
    for line in INSTRUCTIONS:
        ws.append([line])
    ws['A1'].font = Font(bold=True, size=14)
    ws.column_dimensions['A'].width = 110
    edit = PatternFill('solid', fgColor='FFF2CC')
    for title, (data, cols) in tabs.items():
        ws = wb.create_sheet(title)
        ws.append(cols)
        for r in data:
            ws.append([('yes' if r.get(c) is True else '' if r.get(c) is False else r.get(c, '')) for c in cols])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for n, c in enumerate(cols, 1):
            letter = ws.cell(1, n).column_letter
            ws.column_dimensions[letter].width = 60 if 'reasoning' in c else 12 if c in ('id', 'level', 'control') else 24
            if c in ('decision', 'notes'):
                for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=n, max_col=n):
                    row[0].fill = edit
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=False)
        ws.freeze_panes = 'C2'
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def _diagnostics(rows):
    """Why words are not raised: helps any threshold revisit. Counts, not decisions."""
    miss = [r for r in rows if r['outcome'] not in RAISED]
    c = lambda f: sum(1 for r in miss if f(r))
    return ['### Why words are not raised\n',
            f'- Romanization: fewer than 2 of 3 match by sound: {c(lambda r: r["exact"] < 2)}; '
            f'2 match, third differs beyond ee/ii-oo/uu: {c(lambda r: r["exact"] == 2 and r["lvq"] == 0)}',
            f'- Vowelled letters not matched by all 3: {c(lambda r: r["vow_match"] < 3)}',
            f'- Any flag (meaning, dialect, alternative): {c(lambda r: bool(r["flags"]))}',
            f'- Romanization 3/3 and letters 3/3 but blocked by a flag, low confidence or `unsure`: '
            f'{c(lambda r: r["exact"] == 3 and r["vow_match"] == 3)}\n']


if __name__ == '__main__':
    main()
